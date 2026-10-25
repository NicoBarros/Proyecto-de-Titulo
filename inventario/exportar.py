from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Equipo, Insumo

ROJO_INACAP = 'E31C19'
GRIS_OSCURO = '1F2937'
GRIS_TEXTO = '6B7280'

COLORES_ESTADO = {
    Equipo.Estado.DISPONIBLE: ('E3F6E3', '1E6B1E'),
    Equipo.Estado.PRESTADO: ('FEF3C7', '92400E'),
    Equipo.Estado.EN_MANTENIMIENTO: ('F3E8FF', '6B21A8'),
    Equipo.Estado.FUERA_SERVICIO: ('FEE2E2', 'B91C1C'),
}
COLOR_OK = ('E3F6E3', '1E6B1E')
COLOR_ALERTA = ('FEE2E2', 'B91C1C')

BORDE = Border(bottom=Side(style='thin', color='E5E7EB'))
FORMATO_PESOS = '"$"#,##0'


def _escribir(ws, fila, col, valor):
    celda = ws.cell(row=fila, column=col)
    if isinstance(valor, str):
        celda.value = ILLEGAL_CHARACTERS_RE.sub('', valor)
        celda.data_type = 's'
    else:
        celda.value = valor
    return celda


def _pintar(celda, colores, negrita=True):
    fondo, texto = colores
    celda.fill = PatternFill('solid', fgColor=fondo)
    celda.font = Font(color=texto, bold=negrita)
    celda.alignment = Alignment(horizontal='center', vertical='center')


def _encabezado_hoja(ws, titulo, columnas):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
    ws['A1'] = titulo
    ws['A1'].font = Font(size=14, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=ROJO_INACAP)
    ws['A1'].alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columnas)
    ws['A2'] = f'LEICA · INACAP Sede Valdivia · Generado el {timezone.localtime().strftime("%d-%m-%Y %H:%M")}'
    ws['A2'].font = Font(italic=True, color=GRIS_TEXTO)
    ws['A2'].alignment = Alignment(indent=1)


def _tabla(ws, nombre, encabezados, filas, anchos, fila_inicio=4):
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_inicio, column=col, value=texto)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=GRIS_OSCURO)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[fila_inicio].height = 22

    for i, fila in enumerate(filas, start=fila_inicio + 1):
        for col, valor in enumerate(fila, start=1):
            celda = _escribir(ws, i, col, valor)
            celda.border = BORDE
            celda.alignment = Alignment(vertical='center')

    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho

    ultima_fila = fila_inicio + max(len(filas), 1)
    rango = f'A{fila_inicio}:{get_column_letter(len(encabezados))}{ultima_fila}'
    tabla = Table(displayName=nombre, ref=rango)
    tabla.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = ws.cell(row=fila_inicio + 1, column=1)
    return fila_inicio + 1


def _hoja_resumen(ws, equipos, insumos):
    ws.title = 'Resumen'
    _encabezado_hoja(ws, 'Inventario LEICA — Resumen', 3)
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 22

    fila = 4
    ws.cell(row=fila, column=1, value='Equipos por estado').font = Font(bold=True, size=12, color=GRIS_OSCURO)
    fila += 1
    for estado, etiqueta in Equipo.Estado.choices:
        cantidad = sum(1 for e in equipos if e.estado == estado)
        ws.cell(row=fila, column=1, value=etiqueta).border = BORDE
        celda = ws.cell(row=fila, column=2, value=cantidad)
        _pintar(celda, COLORES_ESTADO[estado])
        fila += 1
    ws.cell(row=fila, column=1, value='Total de equipos').font = Font(bold=True)
    ws.cell(row=fila, column=2, value=len(equipos)).font = Font(bold=True)
    ws.cell(row=fila, column=2).alignment = Alignment(horizontal='center')
    fila += 1
    ws.cell(row=fila, column=1, value='Valor total de adquisición').font = Font(bold=True)
    total = sum((e.valor_adquisicion or 0) for e in equipos)
    celda = ws.cell(row=fila, column=2, value=float(total))
    celda.number_format = FORMATO_PESOS
    celda.font = Font(bold=True)

    fila += 2
    ws.cell(row=fila, column=1, value='Insumos').font = Font(bold=True, size=12, color=GRIS_OSCURO)
    fila += 1
    bajo = sum(1 for i in insumos if i.bajo_stock)
    for etiqueta, valor, colores in (
        ('Insumos registrados', len(insumos), None),
        ('Con stock suficiente', len(insumos) - bajo, COLOR_OK),
        ('Con stock bajo o agotado', bajo, COLOR_ALERTA),
    ):
        ws.cell(row=fila, column=1, value=etiqueta).border = BORDE
        celda = ws.cell(row=fila, column=2, value=valor)
        if colores:
            _pintar(celda, colores)
        else:
            celda.alignment = Alignment(horizontal='center')
        fila += 1


def _hoja_equipos(ws, equipos):
    encabezados = ['Código', 'Nombre', 'Modelo', 'Especificaciones', 'Ubicación', 'Estado', 'Valor de adquisición']
    _encabezado_hoja(ws, 'Inventario LEICA — Equipos', len(encabezados))
    filas = [
        [e.codigo_inventario, e.nombre, e.modelo or '', e.especificaciones or '', e.ubicacion or '',
         e.get_estado_display(), float(e.valor_adquisicion) if e.valor_adquisicion is not None else None]
        for e in equipos
    ]
    primera = _tabla(ws, 'Equipos', encabezados, filas, [14, 30, 22, 38, 18, 18, 20])
    for i, e in enumerate(equipos, start=primera):
        ws.cell(row=i, column=1).font = Font(bold=True)
        _pintar(ws.cell(row=i, column=6), COLORES_ESTADO.get(e.estado, COLOR_OK))
        ws.cell(row=i, column=7).number_format = FORMATO_PESOS


def _hoja_insumos(ws, insumos):
    encabezados = ['Nombre', 'Categoría', 'Tipo', 'Disponible', 'Total', 'Umbral mínimo', 'Stock']
    _encabezado_hoja(ws, 'Inventario LEICA — Insumos', len(encabezados))
    filas = [
        [i.nombre, i.categoria or '', i.get_tipo_display(), i.cantidad_disponible, i.cantidad_total,
         i.umbral_minimo, 'Bajo' if i.bajo_stock else 'OK']
        for i in insumos
    ]
    primera = _tabla(ws, 'Insumos', encabezados, filas, [32, 20, 14, 13, 11, 15, 12])
    for fila, i in enumerate(insumos, start=primera):
        ws.cell(row=fila, column=1).font = Font(bold=True)
        for col in (4, 5, 6):
            ws.cell(row=fila, column=col).alignment = Alignment(horizontal='center')
        _pintar(ws.cell(row=fila, column=7), COLOR_ALERTA if i.bajo_stock else COLOR_OK)
        if i.bajo_stock:
            ws.cell(row=fila, column=4).font = Font(bold=True, color=COLOR_ALERTA[1])


def generar_excel_inventario():
    equipos = list(Equipo.objects.order_by('codigo_inventario'))
    insumos = list(Insumo.objects.order_by('nombre'))
    wb = Workbook()
    _hoja_resumen(wb.active, equipos, insumos)
    _hoja_equipos(wb.create_sheet('Equipos'), equipos)
    _hoja_insumos(wb.create_sheet('Insumos'), insumos)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = ROJO_INACAP
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def generar_plantilla_importacion():
    wb = Workbook()
    hojas = (
        (wb.active, 'Equipos', ['Código', 'Nombre', 'Modelo', 'Especificaciones', 'Ubicación', 'Estado', 'Valor de adquisición'],
         [14, 30, 22, 38, 18, 20, 20], 'F', [etiqueta for valor, etiqueta in Equipo.Estado.choices if valor != Equipo.Estado.PRESTADO]),
        (wb.create_sheet(), 'Insumos', ['Nombre', 'Categoría', 'Tipo', 'Total', 'Umbral mínimo'],
         [32, 20, 14, 11, 15], 'C', [etiqueta for _, etiqueta in Insumo.Tipo.choices]),
    )
    for ws, nombre, encabezados, anchos, columna_lista, opciones in hojas:
        ws.title = nombre
        _encabezado_hoja(ws, f'Inventario LEICA — {nombre}', len(encabezados))
        _tabla(ws, nombre, encabezados, [], anchos)
        validacion = DataValidation(type='list', formula1=f'"{",".join(opciones)}"', allow_blank=True)
        ws.add_data_validation(validacion)
        validacion.add(f'{columna_lista}5:{columna_lista}1000')
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = ROJO_INACAP
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
