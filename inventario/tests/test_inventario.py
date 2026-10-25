"""
CRUD de equipos e insumos (inventario) y exportacion a Excel.
"""
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from inventario.models import Equipo, Insumo

from .base import LeicaTestCase


class EquiposTests(LeicaTestCase):

    def setUp(self):
        self.client.force_login(self.encargado)

    def test_listado_y_filtros(self):
        self.crear_equipo('A-1', 'Osciloscopio', ubicacion='Sala 1')
        self.crear_equipo('A-2', 'Multimetro', ubicacion='Sala 2', estado=Equipo.Estado.PRESTADO)
        r = self.client.get('/inventario/', {'q': 'osci'})
        self.assertEqual([e.codigo_inventario for e in r.context['equipos']], ['A-1'])
        r = self.client.get('/inventario/', {'estado': 'prestado'})
        self.assertEqual([e.codigo_inventario for e in r.context['equipos']], ['A-2'])
        r = self.client.get('/inventario/', {'ubicacion': 'sala 1'})
        self.assertEqual(r.context['total_equipos_filtrados'], 1)

    def test_orden_invalido_usa_nombre(self):
        self.crear_equipo('B', 'Zeta')
        self.crear_equipo('A', 'Alfa')
        r = self.client.get('/inventario/', {'orden': 'campo_inexistente'})
        self.assertEqual([e.nombre for e in r.context['equipos']], ['Alfa', 'Zeta'])

    def test_crear_equipo(self):
        r = self.client.post('/inventario/?tab=equipos', {
            'codigo_inventario': 'NEW-1', 'nombre': 'Fuente', 'valor_adquisicion': '500',
        })
        self.assertEqual(r.status_code, 302)
        e = Equipo.objects.get(codigo_inventario='NEW-1')
        self.assertEqual(e.valor_adquisicion, Decimal('500'))

    def test_crear_equipo_codigo_duplicado_no_crea(self):
        self.crear_equipo('DUP-1')
        self.client.post('/inventario/?tab=equipos', {'codigo_inventario': 'DUP-1', 'nombre': 'Otro'})
        self.assertEqual(Equipo.objects.filter(codigo_inventario='DUP-1').count(), 1)

    def test_editar_equipo_incrementa_version(self):
        e = self.crear_equipo()
        r = self.client.post(f'/inventario/equipo/{e.id}/editar/', {'codigo_inventario': e.codigo_inventario, 'nombre': 'Nuevo nombre'})
        self.assertEqual(r.status_code, 302)
        e.refresh_from_db()
        self.assertEqual((e.nombre, e.version), ('Nuevo nombre', 2))

    def test_eliminar_equipo_disponible(self):
        e = self.crear_equipo()
        self.client.post(f'/inventario/equipo/{e.id}/eliminar/')
        self.assertFalse(Equipo.objects.filter(id=e.id).exists())

    def test_no_elimina_equipo_prestado_o_en_mantenimiento(self):
        for estado in (Equipo.Estado.PRESTADO, Equipo.Estado.EN_MANTENIMIENTO):
            with self.subTest(estado=estado):
                e = self.crear_equipo(codigo=f'E-{estado}', estado=estado)
                self.client.post(f'/inventario/equipo/{e.id}/eliminar/')
                self.assertTrue(Equipo.objects.filter(id=e.id).exists())

    def test_eliminar_por_get_no_borra(self):
        e = self.crear_equipo()
        self.client.get(f'/inventario/equipo/{e.id}/eliminar/')
        self.assertTrue(Equipo.objects.filter(id=e.id).exists())


class InsumosTests(LeicaTestCase):

    def setUp(self):
        self.client.force_login(self.calidad)

    def test_crear_insumo_inicializa_disponible(self):
        self.client.post('/inventario/?tab=insumos', {
            'nombre': 'Resistencias', 'tipo': 'perecible', 'cantidad_total': '100', 'umbral_minimo': '10',
        })
        i = Insumo.objects.get(nombre='Resistencias')
        self.assertEqual(i.cantidad_disponible, 100)

    def test_filtro_bajo_stock_y_agotado(self):
        self.crear_insumo('Bajo', cantidad=10, umbral=5).__class__.objects.filter(nombre='Bajo').update(cantidad_disponible=3)
        self.crear_insumo('Agotado', cantidad=10).__class__.objects.filter(nombre='Agotado').update(cantidad_disponible=0)
        self.crear_insumo('Normal', cantidad=10, umbral=1)
        r = self.client.get('/inventario/', {'tab': 'insumos', 'stock': 'bajo'})
        self.assertEqual(sorted(i.nombre for i in r.context['insumos']), ['Agotado', 'Bajo'])
        r = self.client.get('/inventario/', {'tab': 'insumos', 'stock': 'agotado'})
        self.assertEqual([i.nombre for i in r.context['insumos']], ['Agotado'])

    def test_editar_insumo_incrementa_version(self):
        i = self.crear_insumo()
        self.client.post(f'/inventario/insumo/{i.id}/editar/', {
            'nombre': 'Protoboard grande', 'tipo': 'renovable', 'cantidad_total': '10', 'umbral_minimo': '2',
        })
        i.refresh_from_db()
        self.assertEqual((i.nombre, i.version), ('Protoboard grande', 2))

    def test_eliminar_insumo(self):
        i = self.crear_insumo()
        self.client.post(f'/inventario/insumo/{i.id}/eliminar/')
        self.assertFalse(Insumo.objects.filter(id=i.id).exists())


class ExportarTests(LeicaTestCase):

    def exportar(self):
        self.client.force_login(self.encargado)
        r = self.client.get('/inventario/exportar/')
        self.assertEqual(r['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertRegex(r['Content-Disposition'], r'attachment; filename="inventario_leica_\d{8}\.xlsx"')
        return load_workbook(BytesIO(r.content))

    def test_excel_tiene_resumen_equipos_e_insumos(self):
        self.assertEqual(self.exportar().sheetnames, ['Resumen', 'Equipos', 'Insumos'])

    def test_hoja_equipos(self):
        self.crear_equipo('EXP-2', 'Switch', valor=Decimal('250000'), estado=Equipo.Estado.PRESTADO)
        self.crear_equipo('EXP-1', 'Router', valor=None)
        ws = self.exportar()['Equipos']
        self.assertEqual([c.value for c in ws[4]], ['Código', 'Nombre', 'Modelo', 'Especificaciones', 'Ubicación', 'Estado', 'Valor de adquisición'])
        self.assertEqual([ws['A5'].value, ws['F5'].value, ws['G5'].value], ['EXP-1', 'Disponible', None])
        self.assertEqual([ws['A6'].value, ws['F6'].value, ws['G6'].value], ['EXP-2', 'Prestado', 250000])
        self.assertEqual(ws['F5'].fill.fgColor.rgb[-6:], 'E3F6E3')
        self.assertEqual(ws['F6'].fill.fgColor.rgb[-6:], 'FEF3C7')
        self.assertEqual(ws['G6'].number_format, '"$"#,##0')
        self.assertIn('Equipos', ws.tables)

    def test_hoja_insumos_marca_stock_bajo(self):
        self.crear_insumo('Cables', cantidad=10, umbral=5)
        bajo = self.crear_insumo('Resistencias', cantidad=10, umbral=5)
        Insumo.objects.filter(id=bajo.id).update(cantidad_disponible=2)
        ws = self.exportar()['Insumos']
        self.assertEqual([c.value for c in ws[5]], ['Cables', None, 'Renovable', 10, 10, 5, 'OK'])
        self.assertEqual([c.value for c in ws[6]], ['Resistencias', None, 'Renovable', 2, 10, 5, 'Bajo'])
        self.assertEqual(ws['G6'].fill.fgColor.rgb[-6:], 'FEE2E2')

    def test_resumen_cuenta_por_estado(self):
        self.crear_equipo('R-1', valor=Decimal('1000'))
        self.crear_equipo('R-2', valor=Decimal('500'), estado=Equipo.Estado.FUERA_SERVICIO)
        ws = self.exportar()['Resumen']
        valores = {ws.cell(row=f, column=1).value: ws.cell(row=f, column=2).value for f in range(5, ws.max_row + 1)}
        self.assertEqual(valores['Disponible'], 1)
        self.assertEqual(valores['Fuera de servicio'], 1)
        self.assertEqual(valores['Total de equipos'], 2)
        self.assertEqual(valores['Valor total de adquisición'], 1500)

    def test_inventario_vacio_no_falla(self):
        self.assertEqual(self.exportar()['Equipos']['A5'].value, None)

    def test_excel_no_ejecuta_formulas_ni_falla_con_caracteres_de_control(self):
        self.crear_equipo('=1+1', '=HYPERLINK("http://malicioso","clic")', especificaciones='Pegado de Word\x0b')
        ws = self.exportar()['Equipos']
        self.assertEqual((ws['A5'].value, ws['A5'].data_type), ('=1+1', 's'))
        self.assertEqual((ws['B5'].value, ws['B5'].data_type), ('=HYPERLINK("http://malicioso","clic")', 's'))
        self.assertEqual(ws['D5'].value, 'Pegado de Word')


class TrazabilidadTests(LeicaTestCase):

    def test_no_se_elimina_equipo_con_historial(self):
        equipo = self.crear_equipo()
        self.crear_solicitud(self.estudiante, equipos=[equipo], estado='devuelta')
        self.client.force_login(self.encargado)
        self.client.post(f'/inventario/equipo/{equipo.id}/eliminar/')
        self.assertTrue(Equipo.objects.filter(id=equipo.id).exists())

    def test_no_se_elimina_insumo_con_historial(self):
        insumo = self.crear_insumo()
        self.crear_solicitud(self.estudiante, insumos=[(insumo, 1)])
        self.client.force_login(self.encargado)
        self.client.post(f'/inventario/insumo/{insumo.id}/eliminar/')
        self.assertTrue(Insumo.objects.filter(id=insumo.id).exists())
