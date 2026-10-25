from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from inventario.models import Equipo, HistorialEquipo, Insumo, Migracion
from .base import LeicaTestCase


def libro(equipos=None, insumos=None):
    wb = Workbook()
    wb.active.title = 'Notas'
    if equipos is not None:
        ws = wb.create_sheet('Equipos')
        ws.append(['Código', 'Nombre', 'Modelo', 'Especificaciones', 'Ubicación', 'Estado', 'Valor de adquisición'])
        for fila in equipos:
            ws.append(fila)
    if insumos is not None:
        ws = wb.create_sheet('Insumos')
        ws.append(['Nombre', 'Categoría', 'Tipo', 'Total', 'Umbral mínimo'])
        for fila in insumos:
            ws.append(fila)
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


class ImportarInventarioTests(LeicaTestCase):

    def setUp(self):
        self.client.force_login(self.encargado)

    def subir(self, contenido, nombre='inventario.xlsx'):
        return self.client.post('/inventario/importar/', {'archivo': SimpleUploadedFile(nombre, contenido)})

    def test_crea_equipos_e_insumos(self):
        self.subir(libro(
            equipos=[['E-1', 'Osciloscopio', 'TBS1052B', '50 MHz', 'Sala 1', 'Disponible', 450000],
                     ['E-2', 'Multímetro', None, None, None, 'Fuera de servicio', '$35.990']],
            insumos=[['Protoboard', 'Componentes', 'Renovable', 15, 3]],
        ))
        e1, e2 = Equipo.objects.order_by('codigo_inventario')
        self.assertEqual((e1.nombre, e1.valor_adquisicion, e1.estado), ('Osciloscopio', Decimal('450000.00'), 'disponible'))
        self.assertEqual((e2.estado, e2.valor_adquisicion), ('fuera_servicio', Decimal('35990.00')))
        insumo = Insumo.objects.get()
        self.assertEqual((insumo.tipo, insumo.cantidad_total, insumo.cantidad_disponible), ('renovable', 15, 15))
        self.assertEqual(HistorialEquipo.objects.filter(accion=HistorialEquipo.Accion.IMPORTADO).count(), 2)
        m = Migracion.objects.get()
        self.assertEqual((m.filas_totales, m.filas_importadas, m.filas_con_error, m.ejecutado_por),
                         (3, 3, 0, self.encargado))

    def test_actualiza_existentes_sin_borrar_columnas_ausentes(self):
        e = self.crear_equipo('E-1', nombre='Viejo', modelo='M1', ubicacion='Bodega')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Equipos'
        ws.append(['Código', 'Nombre'])
        ws.append(['E-1', 'Nuevo nombre'])
        salida = BytesIO()
        wb.save(salida)
        self.subir(salida.getvalue())
        e.refresh_from_db()
        self.assertEqual((e.nombre, e.modelo, e.ubicacion, e.version), ('Nuevo nombre', 'M1', 'Bodega', 2))
        self.assertTrue(e.historial.filter(accion=HistorialEquipo.Accion.EDITADO).exists())

    def test_no_cambia_el_estado_de_equipos_existentes(self):
        e = self.crear_equipo('E-1', estado=Equipo.Estado.PRESTADO)
        self.subir(libro(equipos=[['E-1', 'Multimetro', None, None, None, 'Disponible', 1000]]))
        e.refresh_from_db()
        self.assertEqual(e.estado, Equipo.Estado.PRESTADO)

    def test_filas_invalidas_se_reportan_y_las_validas_se_importan(self):
        r = self.subir(libro(equipos=[
            ['E-1', 'Válido', None, None, None, None, None],
            ['E-2', None, None, None, None, None, None],
            ['E-3', 'Prestado', None, None, None, 'Prestado', None],
            ['E-4', 'Negativo', None, None, None, None, -5],
        ]))
        self.assertEqual(list(Equipo.objects.values_list('codigo_inventario', flat=True)), ['E-1'])
        m = Migracion.objects.get()
        self.assertEqual((m.filas_importadas, m.filas_con_error), (1, 3))
        self.assertIn('Equipos, fila 3', m.detalle_errores)
        self.assertIn('Equipos, fila 4', m.detalle_errores)
        self.assertRedirects(r, f'/inventario/importar/?detalle={m.id}')
        self.assertContains(self.client.get(r.url), 'Equipos, fila 5')

    def test_insumo_existente_ajusta_disponible(self):
        i = self.crear_insumo('Protoboard', cantidad=10)
        Insumo.objects.filter(id=i.id).update(cantidad_disponible=6)
        self.subir(libro(insumos=[['Protoboard', 'Componentes', 'Renovable', 20, 2], ['Otro', None, 'Raro', 5, 1]]))
        i.refresh_from_db()
        self.assertEqual((i.cantidad_total, i.cantidad_disponible), (20, 16))
        self.assertIn('Tipo', Migracion.objects.get().detalle_errores)

    def test_reimporta_el_archivo_exportado(self):
        self.crear_equipo('E-1', nombre='Osciloscopio')
        self.crear_insumo('Protoboard')
        exportado = self.client.get('/inventario/exportar/').content
        self.subir(exportado)
        m = Migracion.objects.get()
        self.assertEqual((m.filas_totales, m.filas_con_error), (2, 0))
        self.assertEqual(Equipo.objects.count(), 1)

    def test_plantilla_tiene_las_hojas_esperadas(self):
        r = self.client.get('/inventario/importar/plantilla/')
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Equipos', 'Insumos'])
        self.subir(r.content)
        self.assertEqual(Migracion.objects.get().filas_totales, 0)

    def test_rechaza_archivos_que_no_son_excel(self):
        self.subir(b'no es un excel')
        self.subir(b'x', nombre='datos.csv')
        self.subir(libro())
        self.assertFalse(Migracion.objects.exists())

    def test_solo_encargado_y_calidad(self):
        self.client.force_login(self.director)
        self.subir(libro(equipos=[['E-1', 'X', None, None, None, None, None]]))
        self.assertFalse(Equipo.objects.exists())
