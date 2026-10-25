import re
from zipfile import BadZipFile

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from . import auditoria, servicios
from .forms import EquipoForm, InsumoForm
from .models import Equipo, HistorialEquipo, Insumo, Migracion

COLUMNAS_EQUIPOS = {
    'código': 'codigo_inventario',
    'nombre': 'nombre',
    'modelo': 'modelo',
    'especificaciones': 'especificaciones',
    'ubicación': 'ubicacion',
    'estado': 'estado',
    'valor de adquisición': 'valor_adquisicion',
}
COLUMNAS_INSUMOS = {
    'nombre': 'nombre',
    'categoría': 'categoria',
    'tipo': 'tipo',
    'total': 'cantidad_total',
    'umbral mínimo': 'umbral_minimo',
}
ESTADOS_IMPORTABLES = {
    clave: valor
    for valor, etiqueta in Equipo.Estado.choices if valor != Equipo.Estado.PRESTADO
    for clave in (valor, etiqueta.lower())
}
TIPOS_INSUMO = {clave: valor for valor, etiqueta in Insumo.Tipo.choices for clave in (valor, etiqueta.lower())}
DETALLE = 'Importación desde Excel'


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _monto(texto):
    texto = texto.replace('$', '').replace(' ', '')
    if ',' in texto:
        return texto.replace('.', '').replace(',', '.')
    if re.fullmatch(r'\d{1,3}(\.\d{3})+', texto):
        return texto.replace('.', '')
    return texto


def _encabezado(ws, columnas, clave):
    for numero, fila in enumerate(ws.iter_rows(max_row=15, values_only=True), start=1):
        nombres = [_texto(celda).lower() for celda in fila]
        if clave in nombres:
            return numero, {columnas[nombre]: i for i, nombre in enumerate(nombres) if nombre in columnas}
    raise ValueError(f'La hoja "{ws.title}" no tiene una fila de encabezados reconocible.')


def _filas(ws, columnas, clave):
    encabezado, indices = _encabezado(ws, columnas, clave)
    for numero, fila in enumerate(ws.iter_rows(min_row=encabezado + 1, values_only=True), start=encabezado + 1):
        datos = {campo: _texto(fila[i]) if i < len(fila) else '' for campo, i in indices.items()}
        if any(datos.values()):
            yield numero, datos


def _errores_form(form):
    return '; '.join(f"{form.fields[c].label if c in form.fields else c}: {' '.join(m)}" for c, m in form.errors.items())


def _datos_con_actuales(datos, existente, campos):
    base = {campo: getattr(existente, campo) if getattr(existente, campo) is not None else '' for campo in campos} if existente else {}
    base.update(datos)
    return base


def _importar_equipo(datos, usuario):
    estado = datos.pop('estado', '').lower()
    if 'valor_adquisicion' in datos:
        datos['valor_adquisicion'] = _monto(datos['valor_adquisicion'])
    existente = Equipo.objects.select_for_update().filter(codigo_inventario=datos.get('codigo_inventario', '')).first()
    antes = auditoria.foto(existente) if existente else None
    form = EquipoForm(_datos_con_actuales(datos, existente, EquipoForm.Meta.fields), instance=existente)
    if not form.is_valid():
        raise ValueError(_errores_form(form))

    if existente:
        if form.has_changed():
            equipo = form.save(commit=False)
            equipo.version += 1
            equipo.save()
            auditoria.registrar(equipo, usuario, HistorialEquipo.Accion.EDITADO, antes, DETALLE)
        return

    if estado and estado not in ESTADOS_IMPORTABLES:
        raise ValueError(f'Estado: "{estado}" no es válido para un equipo nuevo.')
    equipo = form.save(commit=False)
    equipo.estado = ESTADOS_IMPORTABLES.get(estado, Equipo.Estado.DISPONIBLE)
    equipo.save()
    auditoria.registrar(equipo, usuario, HistorialEquipo.Accion.IMPORTADO, detalle=DETALLE)


def _importar_insumo(datos, usuario):
    if 'tipo' in datos and datos['tipo']:
        tipo = TIPOS_INSUMO.get(datos['tipo'].lower())
        if tipo is None:
            raise ValueError(f'Tipo: "{datos["tipo"]}" no es válido (Renovable o Perecible).')
        datos['tipo'] = tipo
    existente = Insumo.objects.select_for_update().filter(nombre=datos.get('nombre', '')).first()
    en_uso = servicios.unidades_entregadas(existente) if existente else 0
    form = InsumoForm(_datos_con_actuales(datos, existente, InsumoForm.Meta.fields), instance=existente)
    if not form.is_valid():
        raise ValueError(_errores_form(form))

    if existente:
        insumo = form.save(commit=False)
        if insumo.cantidad_total < en_uso:
            raise ValueError(f'Total: hay {en_uso} unidades entregadas; el total no puede ser menor.')
        if form.has_changed():
            servicios.aplicar_total_insumo(insumo, en_uso)
        return

    insumo = form.save(commit=False)
    insumo.cantidad_disponible = insumo.cantidad_total
    insumo.save()


HOJAS = (
    ('equipos', COLUMNAS_EQUIPOS, 'código', _importar_equipo),
    ('insumos', COLUMNAS_INSUMOS, 'nombre', _importar_insumo),
)


def importar_inventario(archivo, usuario):
    try:
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError):
        raise ValueError('El archivo no es un Excel (.xlsx) válido.')

    hojas = {ws.title.strip().lower(): ws for ws in libro.worksheets}
    encontradas = [h for h in HOJAS if h[0] in hojas]
    if not encontradas:
        raise ValueError('El archivo debe tener una hoja llamada "Equipos" o "Insumos".')

    totales = importadas = 0
    errores = []
    with transaction.atomic():
        for nombre, columnas, clave, importar_fila in encontradas:
            for numero, datos in _filas(hojas[nombre], columnas, clave):
                totales += 1
                try:
                    with transaction.atomic():
                        importar_fila(datos, usuario)
                    importadas += 1
                except ValueError as e:
                    errores.append(f'{nombre.capitalize()}, fila {numero}: {e}')

        migracion = Migracion.objects.create(
            archivo_origen=archivo.name[:255],
            filas_totales=totales,
            filas_importadas=importadas,
            filas_con_error=len(errores),
            detalle_errores='\n'.join(errores),
            ejecutado_por=usuario,
        )
    libro.close()
    return migracion
